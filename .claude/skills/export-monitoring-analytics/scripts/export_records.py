#!/usr/bin/env python3
"""Export distinct records and keyword memberships; requires openpyxl."""
import argparse,json,os
from collections import Counter
from datetime import datetime,date,timezone
from pathlib import Path
from zoneinfo import ZoneInfo
from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE

def localtime(value,zone):
 d=datetime.fromisoformat(value.replace('Z','+00:00'))
 if d.tzinfo is None or d.utcoffset() is None:raise ValueError('Timezone-aware dates required')
 return d.astimezone(zone)
def text_cell(ws,row,col,value):
 if isinstance(value,str):
  if len(value)>32767 or ILLEGAL_CHARACTERS_RE.search(value):raise ValueError('Text exceeds Excel limits or contains invalid control characters')
  c=ws.cell(row,col,value);c.data_type='s'
 else:ws.cell(row,col,value)
def sheet(wb,name,rows):
 ws=wb.create_sheet(name)
 for r,values in enumerate(rows,1):
  for c,v in enumerate(values,1):text_cell(ws,r,c,v)
 ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
 for column in ws.columns:
  ws.column_dimensions[column[0].column_letter].width=min(60,max(12,max(len(str(c.value or '')) for c in column)+2))
 return ws

def export(rows,out,field='first_collected',zone_name='UTC',start=None,end=None,scope='All input records; no other filters applied'):
 if not isinstance(rows,list):raise ValueError('Input must be an array')
 if out.suffix.lower()!='.xlsx':raise ValueError('Output must end with .xlsx')
 zone=ZoneInfo(zone_name);lo=date.fromisoformat(start) if start else None;hi=date.fromisoformat(end) if end else None
 if lo and hi and lo>=hi:raise ValueError('start must be before exclusive end')
 seen=set();selected=[];missing=0
 for r in rows:
  if not isinstance(r,dict):raise ValueError('Each record must be an object')
  for k in ('source','id','title'):
   if not isinstance(r.get(k),str):raise ValueError('String required: '+k)
  if not r['source'] or not r['id']:raise ValueError('Empty identity')
  key=(r['source'],r['id'])
  if key in seen:raise ValueError('Duplicate record identity: '+repr(key))
  seen.add(key)
  if not isinstance(r.get('keywords'),list) or any(not isinstance(k,str) for k in r['keywords']):raise ValueError('keywords must be strings')
  value=r.get(field)
  if value is None or value=='':missing+=1;continue
  d=localtime(value,zone)
  if (lo and d.date()<lo) or (hi and d.date()>=hi):continue
  selected.append((r,d))
 sc=Counter(r['source'] for r,d in selected);kc=Counter(k for r,d in selected for k in set(r['keywords']))
 wb=Workbook();wb.remove(wb.active)
 sheet(wb,'Records',[['Source','ID','Title',field+' ('+zone_name+')','Keywords','URL']]+[[r['source'],r['id'],r['title'],d.isoformat(),'; '.join(sorted(set(r['keywords']))),r.get('url','')] for r,d in selected])
 sheet(wb,'Sources',[['Source','Distinct records']]+[[k,v] for k,v in sorted(sc.items())])
 sheet(wb,'Keywords',[['Keyword','Record memberships']]+[[k,v] for k,v in sorted(kc.items())])
 meta={'input_records':len(rows),'exported_records':len(selected),'keyword_memberships':sum(kc.values()),'excluded_missing_date':missing,'excluded_outside_period':len(rows)-missing-len(selected),'date_field':field,'timezone':zone_name,'start_inclusive':start or 'unbounded','end_exclusive':end or 'unbounded','scope':scope,'generated_utc':datetime.now(timezone.utc).isoformat()}
 sheet(wb,'Query',[['Setting','Value']]+list(meta.items()))
 with out.open('xb') as f:
  try:wb.save(f);f.flush();os.fsync(f.fileno())
  except BaseException:
   f.close();out.unlink(missing_ok=True);raise
 return meta

def main():
 p=argparse.ArgumentParser(description=__doc__);p.add_argument('--input',required=True,type=Path);p.add_argument('--output',required=True,type=Path);p.add_argument('--date-field',choices=['first_collected','published'],default='first_collected');p.add_argument('--timezone',default='UTC');p.add_argument('--start');p.add_argument('--end');p.add_argument('--scope',default='All input records; no other filters applied');a=p.parse_args()
 try:print(json.dumps(export(json.loads(a.input.read_text(encoding='utf-8')),a.output,a.date_field,a.timezone,a.start,a.end,a.scope),ensure_ascii=False,indent=2))
 except (ValueError,TypeError,OSError,KeyError) as e:p.exit(2,str(e)+'\n')
if __name__=='__main__':main()
